#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Best-effort importer for Behandlungs_Recherche.xlsx into the LCN MySQL database.

Erweiterung:
- Im DRY RUN werden Preview-Dateien für alle Tabellen erzeugt, in die geschrieben würde.
- Pro Zieltabelle wird mindestens eine CSV-Datei erstellt; zusätzlich gibt es JSON-Dateien
  und eine kleine Manifest-Datei mit Übersicht.

Typical usage:
    python import_behandlungs_recherche_with_previews.py --dry-run
    python import_behandlungs_recherche_with_previews.py --dry-run --preview-dir "C:\\temp\\lcn_preview"
    python import_behandlungs_recherche_with_previews.py

Dependencies:
    pip install pymysql python-dotenv openpyxl
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

pymysql = None
load_dotenv = None
openpyxl = None


def require_pymysql():
    global pymysql
    if pymysql is None:
        try:
            import pymysql as _pymysql
        except ModuleNotFoundError as exc:
            raise ModuleNotFoundError("Modul 'pymysql' fehlt. Bitte installieren mit: pip install pymysql") from exc
        pymysql = _pymysql
    return pymysql


def require_dotenv():
    global load_dotenv
    if load_dotenv is None:
        try:
            from dotenv import load_dotenv as _load_dotenv
        except ModuleNotFoundError as exc:
            raise ModuleNotFoundError("Modul 'python-dotenv' fehlt. Bitte installieren mit: pip install python-dotenv") from exc
        load_dotenv = _load_dotenv
    return load_dotenv


def require_openpyxl():
    global openpyxl
    if openpyxl is None:
        try:
            import openpyxl as _openpyxl
        except ModuleNotFoundError as exc:
            raise ModuleNotFoundError("Modul 'openpyxl' fehlt. Bitte installieren mit: pip install openpyxl") from exc
        openpyxl = _openpyxl
    return openpyxl


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_ENV = SCRIPT_DIR / ".env"
DEFAULT_XLSX = SCRIPT_DIR / "market" / "Behandlungs_Recherche.xlsx"
IMPORT_SOURCE = "Behandlungs_Recherche.xlsx"
STAGING_TABLE = "lcn_import_treatments_market"
BATCH_PREFIX = "behandlungs_recherche"

SOURCE_FIELD_ORDER = [
    "nr",
    "name",
    "typ",
    "einsatzbereich",
    "evidenz_herkunft",
    "community_relevanz",
    "quellen_name",
    "quellen_url",
    "quellen_detail",
    "notiz",
]

MAIN_FIELD_CANDIDATES = {
    "name": ["behandlung", "name", "title", "display_name", "label", "treatment_name", "canonical_name"],
    "slug": ["slug", "name_slug", "url_slug"],
    "sort_order": ["sort_order", "position", "display_order", "order_no", "nr", "rank"],
    "external_id": ["external_id", "source_external_id", "import_external_id", "row_external_id"],
    "import_source": ["import_source", "source_system", "source_name", "import_file", "origin_file"],
    "typ": ["typ", "type", "treatment_type", "category", "kind", "class"],
    "einsatzbereich": [
        "einsatzbereich",
        "indication",
        "indications",
        "use_case",
        "target_area",
        "application_area",
        "symptom_area",
        "indikationen_anwendungsgebiete",
    ],
    "evidenz_herkunft": [
        "evidenz_herkunft",
        "evidence_origin",
        "evidence_source",
        "evidence_basis",
        "evidence_type",
    ],
    "community_relevanz": [
        "community_relevanz",
        "community_priority",
        "community_relevance",
        "relevance",
        "priority",
        "interest_level",
    ],
    "quellen_name": ["quellen_name", "source_name", "reference_name", "citation_name"],
    "quellen_url": ["quellen_url", "source_url", "reference_url", "url"],
    "quellen_detail": ["quellen_detail", "source_detail", "citation_detail", "reference_detail", "details"],
    "notiz": ["weitere_hinweise", "notiz", "note", "notes", "comment", "remarks", "summary", "description"],
}

SOURCE_TABLE_FIELD_CANDIDATES = {
    "name": ["name", "title", "label", "source_name", "reference_name"],
    "url": ["url", "link", "source_url", "reference_url"],
    "detail": ["detail", "details", "citation", "note", "reference_detail", "description"],
    "slug": ["slug", "name_slug", "url_slug"],
    "source_type": ["source_type", "type", "category", "kind"],
    "import_source": ["import_source", "source_system", "origin_file", "import_file"],
    "external_id": ["external_id", "source_external_id", "import_external_id"],
}

ALIAS_TABLE_FIELD_CANDIDATES = {
    "alias": ["alias", "name", "alias_name", "value", "label"],
    "alias_type": ["alias_type", "type", "kind", "alias_kind"],
    "slug": ["slug", "name_slug", "url_slug"],
}

TREATMENT_ID_CANDIDATES = ["treatment_id", "id_treatment", "fk_treatment_id", "treat_id"]
SOURCE_ID_CANDIDATES = ["source_id", "id_source", "fk_source_id"]
ALIAS_TREATMENT_ID_CANDIDATES = ["treatment_id", "id_treatment", "fk_treatment_id", "treat_id"]

SKIP_UPDATE_COLUMNS = {
    "id",
    "created_at",
    "created_by",
    "updated_at",
    "updated_by",
    "deleted_at",
    "password",
}


@dataclass
class TableMatch:
    name: str
    columns: set[str]


@dataclass
class SchemaBundle:
    main_table: Optional[TableMatch]
    alias_table: Optional[TableMatch]
    source_table: Optional[TableMatch]
    treatment_source_link_table: Optional[TableMatch]


# ---------- helpers ----------

def normalize_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, (int, float)) and value == int(value):
        value = str(int(value))
    text = str(value).strip()
    text = re.sub(r"\s+", " ", text)
    return text or None


def slugify(value: Optional[str]) -> Optional[str]:
    value = normalize_text(value)
    if not value:
        return None
    value = unicodedata.normalize("NFKD", value)
    value = value.encode("ascii", "ignore").decode("ascii")
    value = value.lower()
    value = re.sub(r"[^a-z0-9]+", "-", value)
    value = re.sub(r"-{2,}", "-", value).strip("-")
    return value or None


def looks_like_abbreviation(text: str) -> bool:
    t = normalize_text(text) or ""
    if not t:
        return False
    if len(t) > 14:
        return False
    letters = re.sub(r"[^A-Za-z]", "", t)
    return bool(letters) and letters.upper() == letters


def is_descriptor_only(text: str) -> bool:
    t = (normalize_text(text) or "").lower()
    descriptors = {
        "sammelbegriff",
        "oberbegriff",
        "kombination",
        "beispiel",
        "beispiele",
        "gerat",
        "gerät",
        "ansatz",
    }
    return t in descriptors


def truncate(value: Optional[str], max_len: Optional[int]) -> Optional[str]:
    if value is None or max_len is None:
        return value
    if len(value) <= max_len:
        return value
    return value[:max_len]


def print_header(title: str) -> None:
    print("\n" + "=" * 100)
    print(title)
    print("=" * 100)


def sanitize_filename(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9_.-]+", "_", value.strip())
    value = re.sub(r"_+", "_", value).strip("._")
    return value or "preview"


def json_safe(value: Any) -> Any:
    if isinstance(value, (dict, list, str, int, float, bool)) or value is None:
        return value
    return str(value)


def collect_ordered_columns(rows: List[Dict[str, Any]]) -> List[str]:
    cols: List[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row.keys():
            if key not in seen:
                seen.add(key)
                cols.append(key)
    return cols


def add_preview_row(previews: Dict[str, List[Dict[str, Any]]], table_name: str, row: Dict[str, Any]) -> None:
    previews.setdefault(table_name, []).append(row)


def make_preview_ref(table_name: str, existing_id: Optional[Any], payload: Dict[str, Any], fallback_name: Optional[str]) -> str:
    if existing_id is not None:
        return f"{table_name}:id={existing_id}"
    for key in ["slug", "name", "title", "display_name", "label", "external_id", "source_external_id"]:
        if payload.get(key) not in (None, ""):
            return f"{table_name}:{key}={payload[key]}"
    if fallback_name:
        return f"{table_name}:name={fallback_name}"
    return f"{table_name}:unresolved"


def write_preview_files(preview_dir: Path, previews: Dict[str, List[Dict[str, Any]]]) -> List[Tuple[str, int, Path, Path]]:
    preview_dir.mkdir(parents=True, exist_ok=True)
    written: List[Tuple[str, int, Path, Path]] = []

    for table_name, rows in previews.items():
        if not rows:
            continue
        stem = sanitize_filename(table_name)
        csv_path = preview_dir / f"{stem}__preview.csv"
        json_path = preview_dir / f"{stem}__preview.json"
        columns = collect_ordered_columns(rows)

        with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow({c: json_safe(row.get(c)) for c in columns})

        with json_path.open("w", encoding="utf-8") as f:
            json.dump([{k: json_safe(v) for k, v in row.items()} for row in rows], f, ensure_ascii=False, indent=2)

        written.append((table_name, len(rows), csv_path, json_path))

    manifest_path = preview_dir / "_preview_manifest.txt"
    with manifest_path.open("w", encoding="utf-8") as f:
        f.write(f"Preview-Verzeichnis: {preview_dir}\n")
        f.write(f"Erstellt am        : {datetime.now().isoformat(timespec='seconds')}\n\n")
        for table_name, count, csv_path, json_path in written:
            f.write(f"- {table_name}: {count} Zeilen\n")
            f.write(f"  CSV : {csv_path}\n")
            f.write(f"  JSON: {json_path}\n\n")

    return written


# ---------- excel ----------

def read_workbook_rows(xlsx_path: Path) -> List[Dict[str, Optional[str]]]:
    wb = require_openpyxl().load_workbook(xlsx_path, data_only=False)
    ws = wb[wb.sheetnames[0]]

    headers = [normalize_text(c.value) for c in ws[3]]
    if headers != SOURCE_FIELD_ORDER:
        print("WARNUNG: Unerwartete Header-Zeile. Verwende trotzdem die gelesenen Spalten:")
        print(headers)

    result: List[Dict[str, Optional[str]]] = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        raw = dict(zip(headers, row))
        cleaned = {k: normalize_text(v) for k, v in raw.items() if k}
        if not any(cleaned.values()):
            continue
        nr = cleaned.get("nr")
        if nr is None:
            continue
        cleaned["slug"] = slugify(cleaned.get("name"))
        cleaned["aliases_json"] = json.dumps(extract_aliases(cleaned.get("name")), ensure_ascii=False)
        result.append(cleaned)
    return result


# ---------- alias parsing ----------

def extract_aliases(name: Optional[str]) -> List[Dict[str, str]]:
    raw_name = normalize_text(name)
    if not raw_name:
        return []

    aliases: List[Tuple[str, str]] = []

    for part in re.findall(r"\(([^)]{1,40})\)", raw_name):
        p = normalize_text(part)
        if not p or is_descriptor_only(p):
            continue
        if looks_like_abbreviation(p):
            aliases.append((p, "abbreviation"))

    no_paren = re.sub(r"\([^)]*\)", "", raw_name)
    if " / " in no_paren:
        parts = [normalize_text(p) for p in no_paren.split("/")]
        parts = [p for p in parts if p]
        if 1 < len(parts) <= 3:
            for p in parts:
                aliases.append((p, "alternate_name"))

    dedup: List[Dict[str, str]] = []
    seen: set[Tuple[str, str]] = set()
    for alias, alias_type in aliases:
        if alias == raw_name:
            continue
        key = (alias.casefold(), alias_type)
        if key in seen:
            continue
        seen.add(key)
        dedup.append({"alias": alias, "alias_type": alias_type})
    return dedup


# ---------- db ----------

def load_env(env_path: Path) -> Dict[str, Any]:
    if not env_path.exists():
        raise FileNotFoundError(f".env nicht gefunden: {env_path}")

    require_dotenv()(env_path)
    cfg = {
        "host": os.getenv("LCN_DB_HOST", "localhost"),
        "port": int(os.getenv("LCN_DB_PORT", "3306")),
        "user": os.getenv("LCN_DB_USERNAME"),
        "password": os.getenv("LCN_DB_PASSWORD"),
        "database": os.getenv("LCN_DB_DATABASE"),
        "charset": "utf8mb4",
        "cursorclass": require_pymysql().cursors.DictCursor,
        "autocommit": False,
    }
    missing = [k for k in ["user", "password", "database"] if not cfg.get(k)]
    if missing:
        raise RuntimeError(f"Fehlende DB-Werte in .env: {', '.join(missing)}")
    return cfg


def connect_db(cfg: Dict[str, Any]):
    return require_pymysql().connect(**cfg)


def fetchall(cur, sql: str, params: Sequence[Any] | None = None) -> List[Dict[str, Any]]:
    cur.execute(sql, params or [])
    return list(cur.fetchall())


def fetchone(cur, sql: str, params: Sequence[Any] | None = None) -> Optional[Dict[str, Any]]:
    cur.execute(sql, params or [])
    return cur.fetchone()


def get_table_columns(cur, db_name: str) -> Dict[str, List[Dict[str, Any]]]:
    sql = """
        SELECT TABLE_NAME, COLUMN_NAME, DATA_TYPE, COLUMN_TYPE, IS_NULLABLE, CHARACTER_MAXIMUM_LENGTH, COLUMN_KEY
        FROM INFORMATION_SCHEMA.COLUMNS
        WHERE TABLE_SCHEMA = %s
        ORDER BY TABLE_NAME, ORDINAL_POSITION
    """
    rows = fetchall(cur, sql, [db_name])
    result: Dict[str, List[Dict[str, Any]]] = {}
    for row in rows:
        result.setdefault(row["TABLE_NAME"], []).append(row)
    return result


def table_columns_set(col_rows: List[Dict[str, Any]]) -> set[str]:
    return {r["COLUMN_NAME"] for r in col_rows}


def has_any(cols: set[str], names: Iterable[str]) -> bool:
    return any(name in cols for name in names)


def score_table(name: str, cols: set[str], rules: List[Tuple[int, bool]]) -> int:
    score = 0
    for points, ok in rules:
        if ok:
            score += points
    return score


def pick_best_table(table_meta: Dict[str, List[Dict[str, Any]]], scorer) -> Optional[TableMatch]:
    best: Optional[Tuple[int, str, set[str]]] = None
    for table_name, col_rows in table_meta.items():
        cols = table_columns_set(col_rows)
        score = scorer(table_name, cols)
        if score <= 0:
            continue
        if best is None or score > best[0]:
            best = (score, table_name, cols)
    if best is None:
        return None
    return TableMatch(name=best[1], columns=best[2])


def discover_schema(table_meta: Dict[str, List[Dict[str, Any]]]) -> SchemaBundle:
    def pick_exact(name: str) -> Optional[TableMatch]:
        if name not in table_meta:
            return None
        return TableMatch(name=name, columns=table_columns_set(table_meta[name]))

    return SchemaBundle(
        main_table=pick_exact("tbl_treatments_03"),
        alias_table=None,
        source_table=None,
        treatment_source_link_table=None,
    )


def get_id_column(cols: set[str]) -> Optional[str]:
    if "id" in cols:
        return "id"
    for name in ["treat_id", "treatment_id", "source_id", "alias_id", "treatment_alias_id", "treatment_source_id"]:
        if name in cols:
            return name
    return None


def get_colinfo_map(col_rows: List[Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    return {r["COLUMN_NAME"]: r for r in col_rows}


def resolve_existing_column(cols: set[str], candidates: Sequence[str]) -> Optional[str]:
    for c in candidates:
        if c in cols:
            return c
    return None


def ensure_staging_table(cur, dry_run: bool) -> None:
    sql = f"""
        CREATE TABLE IF NOT EXISTS `{STAGING_TABLE}` (
            `id` BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
            `import_batch` VARCHAR(100) NOT NULL,
            `row_nr` INT NOT NULL,
            `name` VARCHAR(255) NULL,
            `typ` VARCHAR(255) NULL,
            `einsatzbereich` TEXT NULL,
            `evidenz_herkunft` TEXT NULL,
            `community_relevanz` VARCHAR(100) NULL,
            `quellen_name` VARCHAR(255) NULL,
            `quellen_url` TEXT NULL,
            `quellen_detail` TEXT NULL,
            `notiz` LONGTEXT NULL,
            `slug` VARCHAR(255) NULL,
            `aliases_json` JSON NULL,
            `created_at` TIMESTAMP NULL DEFAULT CURRENT_TIMESTAMP,
            `updated_at` TIMESTAMP NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
            PRIMARY KEY (`id`),
            UNIQUE KEY `uniq_import_batch_row_nr` (`import_batch`, `row_nr`),
            KEY `idx_name` (`name`),
            KEY `idx_slug` (`slug`)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
    """
    if dry_run:
        print(f"[DRY RUN] CREATE TABLE IF NOT EXISTS `{STAGING_TABLE}` ...")
        return
    cur.execute(sql)


def upsert_staging_row(cur, row: Dict[str, Optional[str]], import_batch: str, dry_run: bool) -> None:
    sql = f"""
        INSERT INTO `{STAGING_TABLE}`
            (`import_batch`, `row_nr`, `name`, `typ`, `einsatzbereich`, `evidenz_herkunft`, `community_relevanz`,
             `quellen_name`, `quellen_url`, `quellen_detail`, `notiz`, `slug`, `aliases_json`)
        VALUES
            (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        ON DUPLICATE KEY UPDATE
            `name` = VALUES(`name`),
            `typ` = VALUES(`typ`),
            `einsatzbereich` = VALUES(`einsatzbereich`),
            `evidenz_herkunft` = VALUES(`evidenz_herkunft`),
            `community_relevanz` = VALUES(`community_relevanz`),
            `quellen_name` = VALUES(`quellen_name`),
            `quellen_url` = VALUES(`quellen_url`),
            `quellen_detail` = VALUES(`quellen_detail`),
            `notiz` = VALUES(`notiz`),
            `slug` = VALUES(`slug`),
            `aliases_json` = VALUES(`aliases_json`)
    """
    params = [
        import_batch,
        int(row["nr"]),
        row.get("name"),
        row.get("typ"),
        row.get("einsatzbereich"),
        row.get("evidenz_herkunft"),
        row.get("community_relevanz"),
        row.get("quellen_name"),
        row.get("quellen_url"),
        row.get("quellen_detail"),
        row.get("notiz"),
        row.get("slug"),
        row.get("aliases_json"),
    ]
    if dry_run:
        print(f"[DRY RUN] staging upsert row_nr={row['nr']} name={row.get('name')!r}")
        return
    cur.execute(sql, params)


def build_insert_dict_for_main(row: Dict[str, Optional[str]], cols: set[str], colinfo: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    payload: Dict[str, Any] = {}

    mapping = {
        "name": row.get("name"),
        "slug": row.get("slug"),
        "typ": row.get("typ"),
        "notiz": row.get("notiz"),
    }

    for logical_key, value in mapping.items():
        if value is None:
            continue
        col = resolve_existing_column(cols, MAIN_FIELD_CANDIDATES[logical_key])
        if not col:
            continue
        max_len = colinfo[col].get("CHARACTER_MAXIMUM_LENGTH")
        if isinstance(value, str):
            value = truncate(value, max_len)
        payload[col] = value

    # Schritt 1: indikationen_anwendungsgebiete bewusst NICHT befüllen
    if "indikationen_anwendungsgebiete" in payload:
        del payload["indikationen_anwendungsgebiete"]

    return payload


def choose_lookup_for_main(cols: set[str], payload: Dict[str, Any]) -> Optional[Tuple[List[str], List[Any]]]:
    slug_col = resolve_existing_column(cols, MAIN_FIELD_CANDIDATES["slug"])
    name_col = resolve_existing_column(cols, MAIN_FIELD_CANDIDATES["name"])

    if slug_col and payload.get(slug_col):
        return ([slug_col], [payload[slug_col]])
    if name_col and payload.get(name_col):
        return ([name_col], [payload[name_col]])
    return None


def select_existing_id(cur, table: str, id_col: str, lookup_cols: List[str], lookup_values: List[Any]) -> Optional[Any]:
    where = " AND ".join([f"`{c}` = %s" for c in lookup_cols])
    sql = f"SELECT `{id_col}` FROM `{table}` WHERE {where} LIMIT 1"
    row = fetchone(cur, sql, lookup_values)
    if not row:
        return None
    return row[id_col]


def insert_row(cur, table: str, payload: Dict[str, Any]) -> int:
    cols = list(payload.keys())
    placeholders = ", ".join(["%s"] * len(cols))
    col_sql = ", ".join([f"`{c}`" for c in cols])
    sql = f"INSERT INTO `{table}` ({col_sql}) VALUES ({placeholders})"
    cur.execute(sql, [payload[c] for c in cols])
    return int(cur.lastrowid)


def update_row(cur, table: str, row_id: Any, id_col: str, payload: Dict[str, Any]) -> None:
    updates = [c for c in payload.keys() if c != id_col and c not in SKIP_UPDATE_COLUMNS]
    if not updates:
        return
    set_sql = ", ".join([f"`{c}` = %s" for c in updates])
    sql = f"UPDATE `{table}` SET {set_sql} WHERE `{id_col}` = %s"
    params = [payload[c] for c in updates] + [row_id]
    cur.execute(sql, params)


def upsert_main_row(
    cur,
    table: TableMatch,
    col_rows: List[Dict[str, Any]],
    row: Dict[str, Optional[str]],
    dry_run: bool,
) -> Tuple[Optional[int], str, Dict[str, Any], Optional[Any]]:
    cols = table.columns
    colinfo = get_colinfo_map(col_rows)
    id_col = get_id_column(cols)
    if not id_col:
        return None, "skipped_table_without_id_column", {}, None

    payload = build_insert_dict_for_main(row, cols, colinfo)

    if not payload:
        return None, "skipped_no_matching_columns", {}, None

    lookup = choose_lookup_for_main(cols, payload)
    existing_id = None
    if lookup:
        lookup_cols, lookup_vals = lookup
        existing_id = select_existing_id(cur, table.name, id_col, lookup_cols, lookup_vals)

    if dry_run:
        action = "update" if existing_id else "insert"
        print(f"[DRY RUN] main {action} {table.name} :: {row.get('name')}")
        return (int(existing_id) if existing_id else None), f"dry_run_{action}", payload, existing_id

    if existing_id:
        update_row(cur, table.name, existing_id, id_col, payload)
        return int(existing_id), "updated", payload, existing_id

    new_id = insert_row(cur, table.name, payload)
    return new_id, "inserted", payload, None


def build_source_payload(row: Dict[str, Optional[str]], cols: set[str], colinfo: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    payload: Dict[str, Any] = {}
    mapping = {
        "name": row.get("quellen_name"),
        "url": row.get("quellen_url"),
        "detail": row.get("quellen_detail"),
        "slug": slugify(row.get("quellen_name")),
        "source_type": "literature",
        "import_source": IMPORT_SOURCE,
        "external_id": row.get("nr"),
    }
    for logical_key, value in mapping.items():
        if value is None:
            continue
        col = resolve_existing_column(cols, SOURCE_TABLE_FIELD_CANDIDATES[logical_key])
        if not col:
            continue
        max_len = colinfo[col].get("CHARACTER_MAXIMUM_LENGTH")
        if isinstance(value, str):
            value = truncate(value, max_len)
        payload[col] = value
    return payload


def choose_lookup_for_source(cols: set[str], payload: Dict[str, Any]) -> Optional[Tuple[List[str], List[Any]]]:
    url_col = resolve_existing_column(cols, SOURCE_TABLE_FIELD_CANDIDATES["url"])
    name_col = resolve_existing_column(cols, SOURCE_TABLE_FIELD_CANDIDATES["name"])
    if url_col and payload.get(url_col):
        return ([url_col], [payload[url_col]])
    if name_col and payload.get(name_col):
        return ([name_col], [payload[name_col]])
    return None


def upsert_source_row(
    cur,
    table: TableMatch,
    col_rows: List[Dict[str, Any]],
    row: Dict[str, Optional[str]],
    dry_run: bool,
) -> Tuple[Optional[int], str, Dict[str, Any], Optional[Any]]:
    if not row.get("quellen_name") and not row.get("quellen_url"):
        return None, "skipped_empty_source", {}, None

    cols = table.columns
    colinfo = get_colinfo_map(col_rows)
    id_col = get_id_column(cols)

    if not id_col:
        return None, "skipped_table_without_id_column", {}, None

    payload = build_source_payload(row, cols, colinfo)

    if not payload:
        return None, "skipped_no_matching_columns", {}, None

    lookup = choose_lookup_for_source(cols, payload)
    existing_id = None
    if lookup:
        lookup_cols, lookup_vals = lookup
        existing_id = select_existing_id(cur, table.name, id_col, lookup_cols, lookup_vals)

    if dry_run:
        action = "update" if existing_id else "insert"
        print(f"[DRY RUN] source {action} {table.name} :: {row.get('quellen_name')}")
        return (int(existing_id) if existing_id else None), f"dry_run_{action}", payload, existing_id

    if existing_id:
        update_row(cur, table.name, existing_id, id_col, payload)
        return int(existing_id), "updated", payload, existing_id

    new_id = insert_row(cur, table.name, payload)
    return new_id, "inserted", payload, None


def link_treatment_source(cur, link_table: TableMatch, treatment_id: int, source_id: int, dry_run: bool) -> str:
    cols = link_table.columns
    tcol = resolve_existing_column(cols, TREATMENT_ID_CANDIDATES)
    scol = resolve_existing_column(cols, SOURCE_ID_CANDIDATES)
    if not tcol or not scol:
        return "skipped_link_columns_missing"

    existing = fetchone(
        cur,
        f"SELECT 1 FROM `{link_table.name}` WHERE `{tcol}` = %s AND `{scol}` = %s LIMIT 1",
        [treatment_id, source_id],
    )
    if existing:
        return "link_exists"

    if dry_run:
        print(f"[DRY RUN] link insert {link_table.name} :: treatment_id={treatment_id} source_id={source_id}")
        return "dry_run_link_insert"

    payload = {tcol: treatment_id, scol: source_id}
    if "is_primary" in cols:
        payload["is_primary"] = 1
    if "sort_order" in cols:
        payload["sort_order"] = 1
    insert_row(cur, link_table.name, payload)
    return "link_inserted"


def upsert_aliases(cur, alias_table: TableMatch, col_rows: List[Dict[str, Any]], treatment_id: int, row: Dict[str, Optional[str]], dry_run: bool) -> Dict[str, int]:
    aliases = json.loads(row.get("aliases_json") or "[]")
    if not aliases:
        return {"alias_inserted": 0, "alias_exists": 0, "alias_skipped": 0}

    cols = alias_table.columns
    colinfo = get_colinfo_map(col_rows)
    alias_col = resolve_existing_column(cols, ALIAS_TABLE_FIELD_CANDIDATES["alias"])
    alias_type_col = resolve_existing_column(cols, ALIAS_TABLE_FIELD_CANDIDATES["alias_type"])
    slug_col = resolve_existing_column(cols, ALIAS_TABLE_FIELD_CANDIDATES["slug"])
    tcol = resolve_existing_column(cols, ALIAS_TREATMENT_ID_CANDIDATES)

    if not alias_col or not tcol:
        return {"alias_inserted": 0, "alias_exists": 0, "alias_skipped": len(aliases)}

    stats = {"alias_inserted": 0, "alias_exists": 0, "alias_skipped": 0}
    for alias_obj in aliases:
        alias = normalize_text(alias_obj.get("alias"))
        alias_type = normalize_text(alias_obj.get("alias_type"))
        if not alias:
            stats["alias_skipped"] += 1
            continue

        existing = fetchone(
            cur,
            f"SELECT 1 AS found FROM `{alias_table.name}` WHERE `{tcol}` = %s AND `{alias_col}` = %s LIMIT 1",
            [treatment_id, alias],
        )
        if existing:
            stats["alias_exists"] += 1
            continue

        payload = {tcol: treatment_id, alias_col: truncate(alias, colinfo[alias_col].get("CHARACTER_MAXIMUM_LENGTH"))}
        if alias_type_col and alias_type:
            payload[alias_type_col] = truncate(alias_type, colinfo[alias_type_col].get("CHARACTER_MAXIMUM_LENGTH"))
        if slug_col:
            payload[slug_col] = truncate(slugify(alias), colinfo[slug_col].get("CHARACTER_MAXIMUM_LENGTH"))

        if dry_run:
            print(f"[DRY RUN] alias insert {alias_table.name} :: treatment_id={treatment_id} alias={alias!r}")
            stats["alias_inserted"] += 1
            continue

        insert_row(cur, alias_table.name, payload)
        stats["alias_inserted"] += 1

    return stats


# ---------- preview builders ----------

def build_staging_preview_row(import_batch: str, row: Dict[str, Optional[str]]) -> Dict[str, Any]:
    return {
        "_preview_action": "upsert",
        "import_batch": import_batch,
        "row_nr": int(row["nr"]),
        "name": row.get("name"),
        "typ": row.get("typ"),
        "einsatzbereich": row.get("einsatzbereich"),
        "evidenz_herkunft": row.get("evidenz_herkunft"),
        "community_relevanz": row.get("community_relevanz"),
        "quellen_name": row.get("quellen_name"),
        "quellen_url": row.get("quellen_url"),
        "quellen_detail": row.get("quellen_detail"),
        "notiz": row.get("notiz"),
        "slug": row.get("slug"),
        "aliases_json": row.get("aliases_json"),
    }


def build_upsert_preview_row(
    row: Dict[str, Optional[str]],
    table_name: str,
    payload: Dict[str, Any],
    state: str,
    existing_id: Optional[Any],
) -> Dict[str, Any]:
    action_map = {
        "inserted": "insert",
        "updated": "update",
        "dry_run_insert": "insert",
        "dry_run_update": "update",
        "skipped_no_matching_columns": "skip",
        "skipped_empty_source": "skip",
        "skipped_table_without_id_column": "skip",
    }
    preview: Dict[str, Any] = {
        "_preview_action": action_map.get(state, state),
        "_row_nr": int(row["nr"]),
        "_input_name": row.get("name"),
    }
    if existing_id is not None:
        preview["_existing_id"] = existing_id
    for key, value in payload.items():
        preview[key] = value
    if not payload:
        preview["_reason"] = state
    return preview


def build_link_preview_row(
    row: Dict[str, Optional[str]],
    link_table: TableMatch,
    treatment_id: Optional[int],
    source_id: Optional[int],
    treatment_ref: Optional[str],
    source_ref: Optional[str],
) -> Dict[str, Any]:
    cols = link_table.columns
    tcol = resolve_existing_column(cols, TREATMENT_ID_CANDIDATES)
    scol = resolve_existing_column(cols, SOURCE_ID_CANDIDATES)

    payload: Dict[str, Any] = {
        "_row_nr": int(row["nr"]),
        "_input_name": row.get("name"),
        "_treatment_ref": treatment_ref,
        "_source_ref": source_ref,
    }
    if tcol:
        payload[tcol] = treatment_id
    if scol:
        payload[scol] = source_id
    if "is_primary" in cols:
        payload["is_primary"] = 1
    if "sort_order" in cols:
        payload["sort_order"] = 1

    if treatment_id is not None and source_id is not None:
        payload["_preview_action"] = "insert_or_exists_checkable"
    else:
        payload["_preview_action"] = "insert_unresolved_ids"
    return payload


def add_alias_preview_rows(
    previews: Dict[str, List[Dict[str, Any]]],
    alias_table: TableMatch,
    alias_col_rows: List[Dict[str, Any]],
    row: Dict[str, Optional[str]],
    treatment_id: Optional[int],
    treatment_ref: Optional[str],
) -> int:
    aliases = json.loads(row.get("aliases_json") or "[]")
    if not aliases:
        return 0

    cols = alias_table.columns
    colinfo = get_colinfo_map(alias_col_rows)
    alias_col = resolve_existing_column(cols, ALIAS_TABLE_FIELD_CANDIDATES["alias"])
    alias_type_col = resolve_existing_column(cols, ALIAS_TABLE_FIELD_CANDIDATES["alias_type"])
    slug_col = resolve_existing_column(cols, ALIAS_TABLE_FIELD_CANDIDATES["slug"])
    tcol = resolve_existing_column(cols, ALIAS_TREATMENT_ID_CANDIDATES)

    if not alias_col or not tcol:
        return 0

    added = 0
    for alias_obj in aliases:
        alias = normalize_text(alias_obj.get("alias"))
        alias_type = normalize_text(alias_obj.get("alias_type"))
        if not alias:
            continue
        rec: Dict[str, Any] = {
            "_row_nr": int(row["nr"]),
            "_input_name": row.get("name"),
            "_treatment_ref": treatment_ref,
            "_preview_action": "insert" if treatment_id is not None else "insert_unresolved_treatment_id",
            tcol: treatment_id,
            alias_col: truncate(alias, colinfo[alias_col].get("CHARACTER_MAXIMUM_LENGTH")),
        }
        if alias_type_col and alias_type:
            rec[alias_type_col] = truncate(alias_type, colinfo[alias_type_col].get("CHARACTER_MAXIMUM_LENGTH"))
        if slug_col:
            rec[slug_col] = truncate(slugify(alias), colinfo[slug_col].get("CHARACTER_MAXIMUM_LENGTH"))
        add_preview_row(previews, alias_table.name, rec)
        added += 1
    return added


# ---------- reporting ----------

def print_schema(bundle: SchemaBundle) -> None:
    print_header("Erkannte Schema-Ziele")
    for label, table in [
        ("Main Treatment Table", bundle.main_table),
        ("Alias Table", bundle.alias_table),
        ("Source Table", bundle.source_table),
        ("Treatment-Source Link Table", bundle.treatment_source_link_table),
    ]:
        if table:
            print(f"{label:28s}: {table.name}")
            print(f"{'':28s}  columns: {', '.join(sorted(table.columns))}")
        else:
            print(f"{label:28s}: NICHT AKTIV IN SCHRITT 1")

# ---------- main ----------

def main() -> int:
    parser = argparse.ArgumentParser(description="Importiert Behandlungs_Recherche.xlsx best-effort in die LCN-Datenbank")
    parser.add_argument("--xlsx", default=str(DEFAULT_XLSX), help="Pfad zur Excel-Datei")
    parser.add_argument("--env", default=str(DEFAULT_ENV), help="Pfad zur .env Datei")
    parser.add_argument("--dry-run", action="store_true", help="Nur anzeigen, was importiert würde")
    parser.add_argument("--limit", type=int, default=None, help="Nur die ersten N Zeilen importieren")
    parser.add_argument("--skip-staging", action="store_true", help="Raw-Staging-Tabelle nicht befüllen")
    parser.add_argument(
        "--preview-dir",
        default=None,
        help="Zielordner für Preview-Dateien. Standard im DRY RUN: <scriptdir>/preview/<import_batch>",
    )
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx)
    env_path = Path(args.env)

    if not xlsx_path.exists():
        print(f"FEHLER: Excel-Datei nicht gefunden: {xlsx_path}")
        return 1

    try:
        rows = read_workbook_rows(xlsx_path)
        if args.limit:
            rows = rows[: args.limit]
        if not rows:
            print("Keine Datensätze gefunden.")
            return 1

        cfg = load_env(env_path)
        batch_ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        import_batch = f"{BATCH_PREFIX}_{batch_ts}"
        preview_dir = Path(args.preview_dir) if args.preview_dir else (SCRIPT_DIR / "preview" / import_batch)

        print_header("Import-Konfiguration")
        print(f"Excel             : {xlsx_path}")
        print(f"ENV               : {env_path}")
        print(f"DB                : {cfg['database']} @ {cfg['host']}:{cfg['port']}")
        print(f"Modus             : {'DRY RUN' if args.dry_run else 'WRITE'}")
        print(f"Datensätze        : {len(rows)}")
        if args.dry_run:
            print(f"Preview-Ordner    : {preview_dir}")

        conn = connect_db(cfg)
        try:
            with conn.cursor() as cur:
                table_meta = get_table_columns(cur, cfg["database"])
                bundle = discover_schema(table_meta)
                print_schema(bundle)

                if not args.skip_staging:
                    ensure_staging_table(cur, args.dry_run)

                stats = {
                    "staging_upserts": 0,
                    "main_inserted": 0,
                    "main_updated": 0,
                    "main_skipped": 0,
                    "source_inserted": 0,
                    "source_updated": 0,
                    "source_skipped": 0,
                    "links_inserted": 0,
                    "links_existing": 0,
                    "links_skipped": 0,
                    "alias_inserted": 0,
                    "alias_exists": 0,
                    "alias_skipped": 0,
                }
                previews: Dict[str, List[Dict[str, Any]]] = {}

                main_col_rows = table_meta.get(bundle.main_table.name, []) if bundle.main_table else []
                alias_col_rows = table_meta.get(bundle.alias_table.name, []) if bundle.alias_table else []
                source_col_rows = table_meta.get(bundle.source_table.name, []) if bundle.source_table else []

                for row in rows:
                    if not args.skip_staging:
                        upsert_staging_row(cur, row, import_batch, args.dry_run)
                        stats["staging_upserts"] += 1
                        if args.dry_run:
                            add_preview_row(previews, STAGING_TABLE, build_staging_preview_row(import_batch, row))

                    treatment_id = None
                    treatment_payload: Dict[str, Any] = {}
                    main_existing_id: Optional[Any] = None
                    treatment_ref: Optional[str] = None
                    main_state = "skipped"

                    if bundle.main_table:
                        treatment_id, main_state, treatment_payload, main_existing_id = upsert_main_row(
                            cur, bundle.main_table, main_col_rows, row, args.dry_run
                        )
                        treatment_ref = make_preview_ref(
                            bundle.main_table.name,
                            main_existing_id if main_existing_id is not None else treatment_id,
                            treatment_payload,
                            row.get("name"),
                        )
                        if main_state in {"inserted", "dry_run_insert"}:
                            stats["main_inserted"] += 1
                        elif main_state in {"updated", "dry_run_update"}:
                            stats["main_updated"] += 1
                        else:
                            stats["main_skipped"] += 1
                        if args.dry_run:
                            add_preview_row(
                                previews,
                                bundle.main_table.name,
                                build_upsert_preview_row(row, bundle.main_table.name, treatment_payload, main_state, main_existing_id),
                            )
                    else:
                        stats["main_skipped"] += 1

                    source_id = None
                    source_payload: Dict[str, Any] = {}
                    source_existing_id: Optional[Any] = None
                    source_ref: Optional[str] = None
                    source_state = "skipped"

                    if bundle.source_table:
                        source_id, source_state, source_payload, source_existing_id = upsert_source_row(
                            cur, bundle.source_table, source_col_rows, row, args.dry_run
                        )
                        source_ref = make_preview_ref(
                            bundle.source_table.name,
                            source_existing_id if source_existing_id is not None else source_id,
                            source_payload,
                            row.get("quellen_name") or row.get("name"),
                        )
                        if source_state in {"inserted", "dry_run_insert"}:
                            stats["source_inserted"] += 1
                        elif source_state in {"updated", "dry_run_update"}:
                            stats["source_updated"] += 1
                        else:
                            stats["source_skipped"] += 1
                        if args.dry_run and source_payload:
                            add_preview_row(
                                previews,
                                bundle.source_table.name,
                                build_upsert_preview_row(row, bundle.source_table.name, source_payload, source_state, source_existing_id),
                            )
                    else:
                        stats["source_skipped"] += 1

                    if bundle.treatment_source_link_table:
                        if args.dry_run:
                            if treatment_payload and source_payload:
                                add_preview_row(
                                    previews,
                                    bundle.treatment_source_link_table.name,
                                    build_link_preview_row(
                                        row,
                                        bundle.treatment_source_link_table,
                                        treatment_id or (int(main_existing_id) if main_existing_id is not None else None),
                                        source_id or (int(source_existing_id) if source_existing_id is not None else None),
                                        treatment_ref,
                                        source_ref,
                                    ),
                                )
                                stats["links_inserted"] += 1
                            else:
                                stats["links_skipped"] += 1
                        else:
                            if treatment_id and source_id:
                                link_state = link_treatment_source(
                                    cur,
                                    bundle.treatment_source_link_table,
                                    treatment_id,
                                    source_id,
                                    args.dry_run,
                                )
                                if link_state in {"link_inserted", "dry_run_link_insert"}:
                                    stats["links_inserted"] += 1
                                elif link_state == "link_exists":
                                    stats["links_existing"] += 1
                                else:
                                    stats["links_skipped"] += 1
                            else:
                                stats["links_skipped"] += 1
                    else:
                        stats["links_skipped"] += 1

                    if bundle.alias_table:
                        if args.dry_run:
                            if treatment_payload:
                                alias_count = add_alias_preview_rows(
                                    previews,
                                    bundle.alias_table,
                                    alias_col_rows,
                                    row,
                                    treatment_id or (int(main_existing_id) if main_existing_id is not None else None),
                                    treatment_ref,
                                )
                                stats["alias_inserted"] += alias_count
                            else:
                                existing_aliases = json.loads(row.get("aliases_json") or "[]")
                                if existing_aliases:
                                    stats["alias_skipped"] += len(existing_aliases)
                        else:
                            if treatment_id:
                                alias_stats = upsert_aliases(cur, bundle.alias_table, alias_col_rows, treatment_id, row, args.dry_run)
                                for k, v in alias_stats.items():
                                    stats[k] += v
                            else:
                                existing_aliases = json.loads(row.get("aliases_json") or "[]")
                                if existing_aliases:
                                    stats["alias_skipped"] += len(existing_aliases)
                    else:
                        existing_aliases = json.loads(row.get("aliases_json") or "[]")
                        if existing_aliases:
                            stats["alias_skipped"] += len(existing_aliases)

                preview_written: List[Tuple[str, int, Path, Path]] = []
                if args.dry_run:
                    preview_written = write_preview_files(preview_dir, previews)
                    conn.rollback()
                else:
                    conn.commit()

                print_header("Import-Ergebnis")
                for key, value in stats.items():
                    print(f"{key:20s}: {value}")

                if args.dry_run:
                    print_header("Preview-Dateien")
                    if preview_written:
                        for table_name, count, csv_path, json_path in preview_written:
                            print(f"{table_name:35s} {count:5d} Zeilen")
                            print(f"  CSV : {csv_path}")
                            print(f"  JSON: {json_path}")
                    else:
                        print("Keine Preview-Dateien erzeugt.")
                    print("\nDRY RUN abgeschlossen. Es wurden keine Änderungen gespeichert.")
                else:
                    print("\nImport abgeschlossen und gespeichert.")

        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()

    except Exception as exc:
        print("\nFEHLER beim Import:")
        print(exc)
        return 1

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
