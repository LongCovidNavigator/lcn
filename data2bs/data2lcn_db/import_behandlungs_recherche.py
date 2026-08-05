#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Best-effort importer for Behandlungs_Recherche.xlsx into the LCN MySQL database.

What the script does:
1) Reads the Excel file (default: ./market/Behandlungs_Recherche.xlsx)
2) Connects to MySQL using the existing .env file
3) Creates/updates a raw staging table so no useful information is lost
4) Introspects the schema and tries to populate the existing treatment-related tables
   as far as the current database structure allows
5) Optionally writes aliases and sources if matching tables exist

Typical usage:
    python import_behandlungs_recherche.py --dry-run
    python import_behandlungs_recherche.py
    python import_behandlungs_recherche.py --xlsx "C:\\xampp\\htdocs\\lcn\\data2bs\\data2lcn_db\\market\\Behandlungs_Recherche.xlsx" --env "C:\\xampp\\htdocs\\lcn\\data2bs\\data2lcn_db\\.env"

Dependencies:
    pip install pymysql python-dotenv openpyxl
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from lcn_env import lcn_env_path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

pymysql = None
load_dotenv = None
openpyxl = None


def require_pymysql():
    global pymysql
    if pymysql is None:
        try:
            import pymysql as _pymysql
        except ModuleNotFoundError as exc:
            raise ModuleNotFoundError("Modul 'pymysql' fehlt. Bitte installieren mit: pip install pymysql") from exc
        pymysql = _pymysql
    return pymysql


def require_dotenv():
    global load_dotenv
    if load_dotenv is None:
        try:
            from dotenv import load_dotenv as _load_dotenv
        except ModuleNotFoundError as exc:
            raise ModuleNotFoundError("Modul 'python-dotenv' fehlt. Bitte installieren mit: pip install python-dotenv") from exc
        load_dotenv = _load_dotenv
    return load_dotenv


def require_openpyxl():
    global openpyxl
    if openpyxl is None:
        try:
            import openpyxl as _openpyxl
        except ModuleNotFoundError as exc:
            raise ModuleNotFoundError("Modul 'openpyxl' fehlt. Bitte installieren mit: pip install openpyxl") from exc
        openpyxl = _openpyxl
    return openpyxl


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_ENV = lcn_env_path()
DEFAULT_XLSX = SCRIPT_DIR / "market" / "Behandlungs_Recherche.xlsx"
IMPORT_SOURCE = "Behandlungs_Recherche.xlsx"
STAGING_TABLE = "lcn_import_treatments_market"
BATCH_PREFIX = "behandlungs_recherche"

SOURCE_FIELD_ORDER = [
    "nr",
    "name",
    "typ",
    "einsatzbereich",
    "evidenz_herkunft",
    "community_relevanz",
    "quellen_name",
    "quellen_url",
    "quellen_detail",
    "notiz",
]

MAIN_FIELD_CANDIDATES = {
    "name": ["name", "title", "display_name", "label", "treatment_name", "canonical_name"],
    "slug": ["slug", "name_slug", "url_slug"],
    "sort_order": ["sort_order", "position", "display_order", "order_no", "nr", "rank"],
    "external_id": ["external_id", "source_external_id", "import_external_id", "row_external_id"],
    "import_source": ["import_source", "source_system", "source_name", "import_file", "origin_file"],
    "typ": ["typ", "type", "treatment_type", "category", "kind", "class"],
    "einsatzbereich": [
        "einsatzbereich",
        "indication",
        "indications",
        "use_case",
        "target_area",
        "application_area",
        "symptom_area",
    ],
    "evidenz_herkunft": [
        "evidenz_herkunft",
        "evidence_origin",
        "evidence_source",
        "evidence_basis",
        "evidence_type",
    ],
    "community_relevanz": [
        "community_relevanz",
        "community_priority",
        "community_relevance",
        "relevance",
        "priority",
        "interest_level",
    ],
    "quellen_name": ["quellen_name", "source_name", "reference_name", "citation_name"],
    "quellen_url": ["quellen_url", "source_url", "reference_url", "url"],
    "quellen_detail": ["quellen_detail", "source_detail", "citation_detail", "reference_detail", "details"],
    "notiz": ["notiz", "note", "notes", "comment", "remarks", "summary", "description"],
}

SOURCE_TABLE_FIELD_CANDIDATES = {
    "name": ["name", "title", "label", "source_name", "reference_name"],
    "url": ["url", "link", "source_url", "reference_url"],
    "detail": ["detail", "details", "citation", "note", "reference_detail", "description"],
    "slug": ["slug", "name_slug", "url_slug"],
    "source_type": ["source_type", "type", "category", "kind"],
    "import_source": ["import_source", "source_system", "origin_file", "import_file"],
    "external_id": ["external_id", "source_external_id", "import_external_id"],
}

ALIAS_TABLE_FIELD_CANDIDATES = {
    "alias": ["alias", "name", "alias_name", "value", "label"],
    "alias_type": ["alias_type", "type", "kind", "alias_kind"],
    "slug": ["slug", "name_slug", "url_slug"],
}

TREATMENT_ID_CANDIDATES = ["treatment_id", "id_treatment", "fk_treatment_id", "treat_id"]
SOURCE_ID_CANDIDATES = ["source_id", "id_source", "fk_source_id"]
ALIAS_TREATMENT_ID_CANDIDATES = ["treatment_id", "id_treatment", "fk_treatment_id", "treat_id"]

SKIP_UPDATE_COLUMNS = {
    "id",
    "created_at",
    "created_by",
    "updated_at",
    "updated_by",
    "deleted_at",
    "password",
}


@dataclass
class TableMatch:
    name: str
    columns: set[str]


@dataclass
class SchemaBundle:
    main_table: Optional[TableMatch]
    alias_table: Optional[TableMatch]
    source_table: Optional[TableMatch]
    treatment_source_link_table: Optional[TableMatch]


# ---------- helpers ----------

def normalize_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, (int, float)) and value == int(value):
        value = str(int(value))
    text = str(value).strip()
    text = re.sub(r"\s+", " ", text)
    return text or None


def slugify(value: Optional[str]) -> Optional[str]:
    value = normalize_text(value)
    if not value:
        return None
    value = unicodedata.normalize("NFKD", value)
    value = value.encode("ascii", "ignore").decode("ascii")
    value = value.lower()
    value = re.sub(r"[^a-z0-9]+", "-", value)
    value = re.sub(r"-{2,}", "-", value).strip("-")
    return value or None


def looks_like_abbreviation(text: str) -> bool:
    t = normalize_text(text) or ""
    if not t:
        return False
    if len(t) > 14:
        return False
    letters = re.sub(r"[^A-Za-z]", "", t)
    return bool(letters) and letters.upper() == letters


def is_descriptor_only(text: str) -> bool:
    t = (normalize_text(text) or "").lower()
    descriptors = {
        "sammelbegriff",
        "oberbegriff",
        "kombination",
        "beispiel",
        "beispiele",
        "gerat",
        "gerät",
        "ansatz",
    }
    return t in descriptors


def truncate(value: Optional[str], max_len: Optional[int]) -> Optional[str]:
    if value is None or max_len is None:
        return value
    if len(value) <= max_len:
        return value
    return value[:max_len]


def print_header(title: str) -> None:
    print("\n" + "=" * 100)
    print(title)
    print("=" * 100)


# ---------- excel ----------

def read_workbook_rows(xlsx_path: Path) -> List[Dict[str, Optional[str]]]:
    wb = require_openpyxl().load_workbook(xlsx_path, data_only=False)
    ws = wb[wb.sheetnames[0]]

    headers = [normalize_text(c.value) for c in ws[3]]
    if headers != SOURCE_FIELD_ORDER:
        print("WARNUNG: Unerwartete Header-Zeile. Verwende trotzdem die gelesenen Spalten:")
        print(headers)

    result: List[Dict[str, Optional[str]]] = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        raw = dict(zip(headers, row))
        cleaned = {k: normalize_text(v) for k, v in raw.items() if k}
        if not any(cleaned.values()):
            continue
        nr = cleaned.get("nr")
        if nr is None:
            continue
        cleaned["slug"] = slugify(cleaned.get("name"))
        cleaned["aliases_json"] = json.dumps(extract_aliases(cleaned.get("name")), ensure_ascii=False)
        result.append(cleaned)
    return result


# ---------- alias parsing ----------

def extract_aliases(name: Optional[str]) -> List[Dict[str, str]]:
    raw_name = normalize_text(name)
    if not raw_name:
        return []

    aliases: List[Tuple[str, str]] = []

    # 1) Acronyms in parentheses, e.g. Low Dose Naltrexone (LDN)
    for part in re.findall(r"\(([^)]{1,40})\)", raw_name):
        p = normalize_text(part)
        if not p or is_descriptor_only(p):
            continue
        if looks_like_abbreviation(p):
            aliases.append((p, "abbreviation"))

    # 2) Slash-separated alternative names, e.g. Pyridostigmin / Mestinon
    no_paren = re.sub(r"\([^)]*\)", "", raw_name)
    if " / " in no_paren:
        parts = [normalize_text(p) for p in no_paren.split("/")]
        parts = [p for p in parts if p]
        if 1 < len(parts) <= 3:
            for p in parts:
                aliases.append((p, "alternate_name"))

    # 3) Clean duplicate aliases / ignore exact raw name
    dedup: List[Dict[str, str]] = []
    seen: set[Tuple[str, str]] = set()
    for alias, alias_type in aliases:
        if alias == raw_name:
            continue
        key = (alias.casefold(), alias_type)
        if key in seen:
            continue
        seen.add(key)
        dedup.append({"alias": alias, "alias_type": alias_type})
    return dedup


# ---------- db ----------

def load_env(env_path: Path) -> Dict[str, Any]:
    if not env_path.exists():
        raise FileNotFoundError(f".env nicht gefunden: {env_path}")

    require_dotenv()(env_path)
    cfg = {
        "host": os.getenv("LCN_DB_HOST", "localhost"),
        "port": int(os.getenv("LCN_DB_PORT", "3306")),
        "user": os.getenv("LCN_DB_USERNAME"),
        "password": os.getenv("LCN_DB_PASSWORD"),
        "database": os.getenv("LCN_DB_DATABASE"),
        "charset": "utf8mb4",
        "cursorclass": require_pymysql().cursors.DictCursor,
        "autocommit": False,
    }
    missing = [k for k in ["user", "password", "database"] if not cfg.get(k)]
    if missing:
        raise RuntimeError(f"Fehlende DB-Werte in .env: {', '.join(missing)}")
    return cfg


def connect_db(cfg: Dict[str, Any]):
    return require_pymysql().connect(**cfg)


def fetchall(cur, sql: str, params: Sequence[Any] | None = None) -> List[Dict[str, Any]]:
    cur.execute(sql, params or [])
    return list(cur.fetchall())


def fetchone(cur, sql: str, params: Sequence[Any] | None = None) -> Optional[Dict[str, Any]]:
    cur.execute(sql, params or [])
    return cur.fetchone()


def get_table_columns(cur, db_name: str) -> Dict[str, List[Dict[str, Any]]]:
    sql = """
        SELECT TABLE_NAME, COLUMN_NAME, DATA_TYPE, COLUMN_TYPE, IS_NULLABLE, CHARACTER_MAXIMUM_LENGTH, COLUMN_KEY
        FROM INFORMATION_SCHEMA.COLUMNS
        WHERE TABLE_SCHEMA = %s
        ORDER BY TABLE_NAME, ORDINAL_POSITION
    """
    rows = fetchall(cur, sql, [db_name])
    result: Dict[str, List[Dict[str, Any]]] = {}
    for row in rows:
        result.setdefault(row["TABLE_NAME"], []).append(row)
    return result


def table_columns_set(col_rows: List[Dict[str, Any]]) -> set[str]:
    return {r["COLUMN_NAME"] for r in col_rows}


def has_any(cols: set[str], names: Iterable[str]) -> bool:
    return any(name in cols for name in names)


def score_table(name: str, cols: set[str], rules: List[Tuple[int, bool]]) -> int:
    score = 0
    for points, ok in rules:
        if ok:
            score += points
    return score


def pick_best_table(table_meta: Dict[str, List[Dict[str, Any]]], scorer) -> Optional[TableMatch]:
    best: Optional[Tuple[int, str, set[str]]] = None
    for table_name, col_rows in table_meta.items():
        cols = table_columns_set(col_rows)
        score = scorer(table_name, cols)
        if score <= 0:
            continue
        if best is None or score > best[0]:
            best = (score, table_name, cols)
    if best is None:
        return None
    return TableMatch(name=best[1], columns=best[2])


def discover_schema(table_meta: Dict[str, List[Dict[str, Any]]]) -> SchemaBundle:
    def score_main(table_name: str, cols: set[str]) -> int:
        lname = table_name.lower()
        rules = [
            (60, "treat" in lname or "behand" in lname),
            (20, has_any(cols, MAIN_FIELD_CANDIDATES["name"])),
            (10, has_any(cols, MAIN_FIELD_CANDIDATES["typ"])),
            (10, has_any(cols, MAIN_FIELD_CANDIDATES["einsatzbereich"])),
            (8, has_any(cols, MAIN_FIELD_CANDIDATES["notiz"])),
            (5, "id" in cols),
        ]
        if "alias" in lname or "source" in lname or "link" in lname:
            return 0
        return score_table(table_name, cols, rules)

    def score_alias(table_name: str, cols: set[str]) -> int:
        lname = table_name.lower()
        rules = [
            (50, "alias" in lname),
            (20, has_any(cols, ALIAS_TREATMENT_ID_CANDIDATES)),
            (20, has_any(cols, ALIAS_TABLE_FIELD_CANDIDATES["alias"])),
            (5, has_any(cols, ALIAS_TABLE_FIELD_CANDIDATES["alias_type"])),
        ]
        return score_table(table_name, cols, rules)

    def score_source(table_name: str, cols: set[str]) -> int:
        lname = table_name.lower()
        rules = [
            (45, "source" in lname or "refer" in lname or "citation" in lname),
            (20, has_any(cols, SOURCE_TABLE_FIELD_CANDIDATES["name"])),
            (20, has_any(cols, SOURCE_TABLE_FIELD_CANDIDATES["url"])),
            (5, "id" in cols),
        ]
        if "link" in lname or (has_any(cols, TREATMENT_ID_CANDIDATES) and has_any(cols, SOURCE_ID_CANDIDATES)):
            return 0
        return score_table(table_name, cols, rules)

    def score_treatment_source_link(table_name: str, cols: set[str]) -> int:
        lname = table_name.lower()
        rules = [
            (50, has_any(cols, TREATMENT_ID_CANDIDATES) and has_any(cols, SOURCE_ID_CANDIDATES)),
            (15, "link" in lname or "cpl" in lname or "map" in lname),
            (5, "source" in lname and ("treat" in lname or "behand" in lname)),
        ]
        return score_table(table_name, cols, rules)

    return SchemaBundle(
        main_table=pick_best_table(table_meta, score_main),
        alias_table=pick_best_table(table_meta, score_alias),
        source_table=pick_best_table(table_meta, score_source),
        treatment_source_link_table=pick_best_table(table_meta, score_treatment_source_link),
    )


def get_id_column(cols: set[str]) -> str:
    if "id" in cols:
        return "id"
    for name in ["treatment_id", "source_id", "alias_id"]:
        if name in cols:
            return name
    raise RuntimeError(f"Keine ID-Spalte gefunden in: {sorted(cols)}")


def get_colinfo_map(col_rows: List[Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    return {r["COLUMN_NAME"]: r for r in col_rows}


def resolve_existing_column(cols: set[str], candidates: Sequence[str]) -> Optional[str]:
    for c in candidates:
        if c in cols:
            return c
    return None


def ensure_staging_table(cur, dry_run: bool) -> None:
    sql = f"""
        CREATE TABLE IF NOT EXISTS `{STAGING_TABLE}` (
            `id` BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
            `import_batch` VARCHAR(100) NOT NULL,
            `row_nr` INT NOT NULL,
            `name` VARCHAR(255) NULL,
            `typ` VARCHAR(255) NULL,
            `einsatzbereich` TEXT NULL,
            `evidenz_herkunft` TEXT NULL,
            `community_relevanz` VARCHAR(100) NULL,
            `quellen_name` VARCHAR(255) NULL,
            `quellen_url` TEXT NULL,
            `quellen_detail` TEXT NULL,
            `notiz` LONGTEXT NULL,
            `slug` VARCHAR(255) NULL,
            `aliases_json` JSON NULL,
            `created_at` TIMESTAMP NULL DEFAULT CURRENT_TIMESTAMP,
            `updated_at` TIMESTAMP NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
            PRIMARY KEY (`id`),
            UNIQUE KEY `uniq_import_batch_row_nr` (`import_batch`, `row_nr`),
            KEY `idx_name` (`name`),
            KEY `idx_slug` (`slug`)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
    """
    if dry_run:
        print(f"[DRY RUN] CREATE TABLE IF NOT EXISTS `{STAGING_TABLE}` ...")
        return
    cur.execute(sql)


def upsert_staging_row(cur, row: Dict[str, Optional[str]], import_batch: str, dry_run: bool) -> None:
    sql = f"""
        INSERT INTO `{STAGING_TABLE}`
            (`import_batch`, `row_nr`, `name`, `typ`, `einsatzbereich`, `evidenz_herkunft`, `community_relevanz`,
             `quellen_name`, `quellen_url`, `quellen_detail`, `notiz`, `slug`, `aliases_json`)
        VALUES
            (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        ON DUPLICATE KEY UPDATE
            `name` = VALUES(`name`),
            `typ` = VALUES(`typ`),
            `einsatzbereich` = VALUES(`einsatzbereich`),
            `evidenz_herkunft` = VALUES(`evidenz_herkunft`),
            `community_relevanz` = VALUES(`community_relevanz`),
            `quellen_name` = VALUES(`quellen_name`),
            `quellen_url` = VALUES(`quellen_url`),
            `quellen_detail` = VALUES(`quellen_detail`),
            `notiz` = VALUES(`notiz`),
            `slug` = VALUES(`slug`),
            `aliases_json` = VALUES(`aliases_json`)
    """
    params = [
        import_batch,
        int(row["nr"]),
        row.get("name"),
        row.get("typ"),
        row.get("einsatzbereich"),
        row.get("evidenz_herkunft"),
        row.get("community_relevanz"),
        row.get("quellen_name"),
        row.get("quellen_url"),
        row.get("quellen_detail"),
        row.get("notiz"),
        row.get("slug"),
        row.get("aliases_json"),
    ]
    if dry_run:
        print(f"[DRY RUN] staging upsert row_nr={row['nr']} name={row.get('name')!r}")
        return
    cur.execute(sql, params)


def build_insert_dict_for_main(row: Dict[str, Optional[str]], cols: set[str], colinfo: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    payload: Dict[str, Any] = {}

    mapping = {
        "name": row.get("name"),
        "slug": row.get("slug"),
        "sort_order": int(row["nr"]) if row.get("nr") else None,
        "external_id": row.get("nr"),
        "import_source": IMPORT_SOURCE,
        "typ": row.get("typ"),
        "einsatzbereich": row.get("einsatzbereich"),
        "evidenz_herkunft": row.get("evidenz_herkunft"),
        "community_relevanz": row.get("community_relevanz"),
        "quellen_name": row.get("quellen_name"),
        "quellen_url": row.get("quellen_url"),
        "quellen_detail": row.get("quellen_detail"),
        "notiz": row.get("notiz"),
    }

    for logical_key, value in mapping.items():
        if value is None:
            continue
        col = resolve_existing_column(cols, MAIN_FIELD_CANDIDATES[logical_key])
        if not col:
            continue
        max_len = colinfo[col].get("CHARACTER_MAXIMUM_LENGTH")
        if isinstance(value, str):
            value = truncate(value, max_len)
        payload[col] = value

    # useful extras if present
    if "is_active" in cols and "is_active" not in payload:
        payload["is_active"] = 1
    if "active" in cols and "active" not in payload:
        payload["active"] = 1

    return payload


def choose_lookup_for_main(cols: set[str], payload: Dict[str, Any]) -> Optional[Tuple[List[str], List[Any]]]:
    ext_col = resolve_existing_column(cols, MAIN_FIELD_CANDIDATES["external_id"])
    src_col = resolve_existing_column(cols, MAIN_FIELD_CANDIDATES["import_source"])
    name_col = resolve_existing_column(cols, MAIN_FIELD_CANDIDATES["name"])
    slug_col = resolve_existing_column(cols, MAIN_FIELD_CANDIDATES["slug"])

    if ext_col and src_col and payload.get(ext_col) and payload.get(src_col):
        return ([ext_col, src_col], [payload[ext_col], payload[src_col]])
    if slug_col and payload.get(slug_col):
        return ([slug_col], [payload[slug_col]])
    if name_col and payload.get(name_col):
        return ([name_col], [payload[name_col]])
    return None


def select_existing_id(cur, table: str, id_col: str, lookup_cols: List[str], lookup_values: List[Any]) -> Optional[Any]:
    where = " AND ".join([f"`{c}` = %s" for c in lookup_cols])
    sql = f"SELECT `{id_col}` FROM `{table}` WHERE {where} LIMIT 1"
    row = fetchone(cur, sql, lookup_values)
    if not row:
        return None
    return row[id_col]


def insert_row(cur, table: str, payload: Dict[str, Any]) -> int:
    cols = list(payload.keys())
    placeholders = ", ".join(["%s"] * len(cols))
    col_sql = ", ".join([f"`{c}`" for c in cols])
    sql = f"INSERT INTO `{table}` ({col_sql}) VALUES ({placeholders})"
    cur.execute(sql, [payload[c] for c in cols])
    return int(cur.lastrowid)


def update_row(cur, table: str, row_id: Any, id_col: str, payload: Dict[str, Any]) -> None:
    updates = [c for c in payload.keys() if c != id_col and c not in SKIP_UPDATE_COLUMNS]
    if not updates:
        return
    set_sql = ", ".join([f"`{c}` = %s" for c in updates])
    sql = f"UPDATE `{table}` SET {set_sql} WHERE `{id_col}` = %s"
    params = [payload[c] for c in updates] + [row_id]
    cur.execute(sql, params)


def upsert_main_row(cur, table: TableMatch, col_rows: List[Dict[str, Any]], row: Dict[str, Optional[str]], dry_run: bool) -> Tuple[Optional[int], str]:
    cols = table.columns
    colinfo = get_colinfo_map(col_rows)
    id_col = get_id_column(cols)
    payload = build_insert_dict_for_main(row, cols, colinfo)
    if not payload:
        return None, "skipped_no_matching_columns"

    lookup = choose_lookup_for_main(cols, payload)
    existing_id = None
    if lookup:
        lookup_cols, lookup_vals = lookup
        existing_id = select_existing_id(cur, table.name, id_col, lookup_cols, lookup_vals)

    if dry_run:
        action = "update" if existing_id else "insert"
        print(f"[DRY RUN] main {action} {table.name} :: {row.get('name')}")
        return (int(existing_id) if existing_id else None), f"dry_run_{action}"

    if existing_id:
        update_row(cur, table.name, existing_id, id_col, payload)
        return int(existing_id), "updated"

    new_id = insert_row(cur, table.name, payload)
    return new_id, "inserted"


def build_source_payload(row: Dict[str, Optional[str]], cols: set[str], colinfo: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    payload: Dict[str, Any] = {}
    mapping = {
        "name": row.get("quellen_name"),
        "url": row.get("quellen_url"),
        "detail": row.get("quellen_detail"),
        "slug": slugify(row.get("quellen_name")),
        "source_type": "literature",
        "import_source": IMPORT_SOURCE,
        "external_id": row.get("nr"),
    }
    for logical_key, value in mapping.items():
        if value is None:
            continue
        col = resolve_existing_column(cols, SOURCE_TABLE_FIELD_CANDIDATES[logical_key])
        if not col:
            continue
        max_len = colinfo[col].get("CHARACTER_MAXIMUM_LENGTH")
        if isinstance(value, str):
            value = truncate(value, max_len)
        payload[col] = value
    return payload


def choose_lookup_for_source(cols: set[str], payload: Dict[str, Any]) -> Optional[Tuple[List[str], List[Any]]]:
    url_col = resolve_existing_column(cols, SOURCE_TABLE_FIELD_CANDIDATES["url"])
    name_col = resolve_existing_column(cols, SOURCE_TABLE_FIELD_CANDIDATES["name"])
    if url_col and payload.get(url_col):
        return ([url_col], [payload[url_col]])
    if name_col and payload.get(name_col):
        return ([name_col], [payload[name_col]])
    return None


def upsert_source_row(cur, table: TableMatch, col_rows: List[Dict[str, Any]], row: Dict[str, Optional[str]], dry_run: bool) -> Tuple[Optional[int], str]:
    if not row.get("quellen_name") and not row.get("quellen_url"):
        return None, "skipped_empty_source"

    cols = table.columns
    colinfo = get_colinfo_map(col_rows)
    id_col = get_id_column(cols)
    payload = build_source_payload(row, cols, colinfo)
    if not payload:
        return None, "skipped_no_matching_columns"

    lookup = choose_lookup_for_source(cols, payload)
    existing_id = None
    if lookup:
        lookup_cols, lookup_vals = lookup
        existing_id = select_existing_id(cur, table.name, id_col, lookup_cols, lookup_vals)

    if dry_run:
        action = "update" if existing_id else "insert"
        print(f"[DRY RUN] source {action} {table.name} :: {row.get('quellen_name')}")
        return (int(existing_id) if existing_id else None), f"dry_run_{action}"

    if existing_id:
        update_row(cur, table.name, existing_id, id_col, payload)
        return int(existing_id), "updated"

    new_id = insert_row(cur, table.name, payload)
    return new_id, "inserted"


def link_treatment_source(cur, link_table: TableMatch, treatment_id: int, source_id: int, dry_run: bool) -> str:
    cols = link_table.columns
    tcol = resolve_existing_column(cols, TREATMENT_ID_CANDIDATES)
    scol = resolve_existing_column(cols, SOURCE_ID_CANDIDATES)
    if not tcol or not scol:
        return "skipped_link_columns_missing"

    existing = fetchone(
        cur,
        f"SELECT 1 FROM `{link_table.name}` WHERE `{tcol}` = %s AND `{scol}` = %s LIMIT 1",
        [treatment_id, source_id],
    )
    if existing:
        return "link_exists"

    if dry_run:
        print(f"[DRY RUN] link insert {link_table.name} :: treatment_id={treatment_id} source_id={source_id}")
        return "dry_run_link_insert"

    payload = {tcol: treatment_id, scol: source_id}
    if "is_primary" in cols:
        payload["is_primary"] = 1
    if "sort_order" in cols:
        payload["sort_order"] = 1
    insert_row(cur, link_table.name, payload)
    return "link_inserted"


def upsert_aliases(cur, alias_table: TableMatch, col_rows: List[Dict[str, Any]], treatment_id: int, row: Dict[str, Optional[str]], dry_run: bool) -> Dict[str, int]:
    aliases = json.loads(row.get("aliases_json") or "[]")
    if not aliases:
        return {"alias_inserted": 0, "alias_exists": 0, "alias_skipped": 0}

    cols = alias_table.columns
    colinfo = get_colinfo_map(col_rows)
    id_col = get_id_column(cols)
    alias_col = resolve_existing_column(cols, ALIAS_TABLE_FIELD_CANDIDATES["alias"])
    alias_type_col = resolve_existing_column(cols, ALIAS_TABLE_FIELD_CANDIDATES["alias_type"])
    slug_col = resolve_existing_column(cols, ALIAS_TABLE_FIELD_CANDIDATES["slug"])
    tcol = resolve_existing_column(cols, ALIAS_TREATMENT_ID_CANDIDATES)

    if not alias_col or not tcol:
        return {"alias_inserted": 0, "alias_exists": 0, "alias_skipped": len(aliases)}

    stats = {"alias_inserted": 0, "alias_exists": 0, "alias_skipped": 0}
    for alias_obj in aliases:
        alias = normalize_text(alias_obj.get("alias"))
        alias_type = normalize_text(alias_obj.get("alias_type"))
        if not alias:
            stats["alias_skipped"] += 1
            continue

        existing = fetchone(
            cur,
            f"SELECT `{id_col}` FROM `{alias_table.name}` WHERE `{tcol}` = %s AND `{alias_col}` = %s LIMIT 1",
            [treatment_id, alias],
        )
        if existing:
            stats["alias_exists"] += 1
            continue

        payload = {tcol: treatment_id, alias_col: truncate(alias, colinfo[alias_col].get("CHARACTER_MAXIMUM_LENGTH"))}
        if alias_type_col and alias_type:
            payload[alias_type_col] = truncate(alias_type, colinfo[alias_type_col].get("CHARACTER_MAXIMUM_LENGTH"))
        if slug_col:
            payload[slug_col] = truncate(slugify(alias), colinfo[slug_col].get("CHARACTER_MAXIMUM_LENGTH"))

        if dry_run:
            print(f"[DRY RUN] alias insert {alias_table.name} :: treatment_id={treatment_id} alias={alias!r}")
            stats["alias_inserted"] += 1
            continue

        insert_row(cur, alias_table.name, payload)
        stats["alias_inserted"] += 1

    return stats


# ---------- reporting ----------

def print_schema(bundle: SchemaBundle) -> None:
    print_header("Erkannte Schema-Ziele")
    for label, table in [
        ("Main Treatment Table", bundle.main_table),
        ("Alias Table", bundle.alias_table),
        ("Source Table", bundle.source_table),
        ("Treatment-Source Link Table", bundle.treatment_source_link_table),
    ]:
        if table:
            print(f"{label:28s}: {table.name}")
        else:
            print(f"{label:28s}: NICHT GEFUNDEN")


# ---------- main ----------

def main() -> int:
    parser = argparse.ArgumentParser(description="Importiert Behandlungs_Recherche.xlsx best-effort in die LCN-Datenbank")
    parser.add_argument("--xlsx", default=str(DEFAULT_XLSX), help="Pfad zur Excel-Datei")
    parser.add_argument("--env", default=str(DEFAULT_ENV), help="Pfad zur .env Datei")
    parser.add_argument("--dry-run", action="store_true", help="Nur anzeigen, was importiert würde")
    parser.add_argument("--limit", type=int, default=None, help="Nur die ersten N Zeilen importieren")
    parser.add_argument("--skip-staging", action="store_true", help="Raw-Staging-Tabelle nicht befüllen")
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx)
    env_path = Path(args.env)

    if not xlsx_path.exists():
        print(f"FEHLER: Excel-Datei nicht gefunden: {xlsx_path}")
        return 1

    try:
        rows = read_workbook_rows(xlsx_path)
        if args.limit:
            rows = rows[: args.limit]
        if not rows:
            print("Keine Datensätze gefunden.")
            return 1

        cfg = load_env(env_path)
        print_header("Import-Konfiguration")
        print(f"Excel             : {xlsx_path}")
        print(f"ENV               : {env_path}")
        print(f"DB                : {cfg['database']} @ {cfg['host']}:{cfg['port']}")
        print(f"Modus             : {'DRY RUN' if args.dry_run else 'WRITE'}")
        print(f"Datensätze        : {len(rows)}")

        conn = connect_db(cfg)
        try:
            with conn.cursor() as cur:
                table_meta = get_table_columns(cur, cfg["database"])
                bundle = discover_schema(table_meta)
                print_schema(bundle)

                if not args.skip_staging:
                    ensure_staging_table(cur, args.dry_run)

                batch_ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                import_batch = f"{BATCH_PREFIX}_{batch_ts}"

                stats = {
                    "staging_upserts": 0,
                    "main_inserted": 0,
                    "main_updated": 0,
                    "main_skipped": 0,
                    "source_inserted": 0,
                    "source_updated": 0,
                    "source_skipped": 0,
                    "links_inserted": 0,
                    "links_existing": 0,
                    "links_skipped": 0,
                    "alias_inserted": 0,
                    "alias_exists": 0,
                    "alias_skipped": 0,
                }

                main_col_rows = table_meta.get(bundle.main_table.name, []) if bundle.main_table else []
                alias_col_rows = table_meta.get(bundle.alias_table.name, []) if bundle.alias_table else []
                source_col_rows = table_meta.get(bundle.source_table.name, []) if bundle.source_table else []

                for row in rows:
                    if not args.skip_staging:
                        upsert_staging_row(cur, row, import_batch, args.dry_run)
                        stats["staging_upserts"] += 1

                    treatment_id = None
                    if bundle.main_table:
                        treatment_id, main_state = upsert_main_row(cur, bundle.main_table, main_col_rows, row, args.dry_run)
                        if main_state in {"inserted", "dry_run_insert"}:
                            stats["main_inserted"] += 1
                        elif main_state in {"updated", "dry_run_update"}:
                            stats["main_updated"] += 1
                        else:
                            stats["main_skipped"] += 1
                    else:
                        stats["main_skipped"] += 1

                    source_id = None
                    if bundle.source_table:
                        source_id, source_state = upsert_source_row(cur, bundle.source_table, source_col_rows, row, args.dry_run)
                        if source_state in {"inserted", "dry_run_insert"}:
                            stats["source_inserted"] += 1
                        elif source_state in {"updated", "dry_run_update"}:
                            stats["source_updated"] += 1
                        else:
                            stats["source_skipped"] += 1
                    else:
                        stats["source_skipped"] += 1

                    if treatment_id and source_id and bundle.treatment_source_link_table:
                        link_state = link_treatment_source(cur, bundle.treatment_source_link_table, treatment_id, source_id, args.dry_run)
                        if link_state in {"link_inserted", "dry_run_link_insert"}:
                            stats["links_inserted"] += 1
                        elif link_state == "link_exists":
                            stats["links_existing"] += 1
                        else:
                            stats["links_skipped"] += 1
                    else:
                        stats["links_skipped"] += 1

                    if treatment_id and bundle.alias_table:
                        alias_stats = upsert_aliases(cur, bundle.alias_table, alias_col_rows, treatment_id, row, args.dry_run)
                        for k, v in alias_stats.items():
                            stats[k] += v
                    else:
                        # count aliases as skipped only if there actually were aliases
                        existing_aliases = json.loads(row.get("aliases_json") or "[]")
                        if existing_aliases:
                            stats["alias_skipped"] += len(existing_aliases)

                if args.dry_run:
                    conn.rollback()
                else:
                    conn.commit()

                print_header("Import-Ergebnis")
                for key, value in stats.items():
                    print(f"{key:20s}: {value}")

                if args.dry_run:
                    print("\nDRY RUN abgeschlossen. Es wurden keine Änderungen gespeichert.")
                else:
                    print("\nImport abgeschlossen und gespeichert.")

        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()

    except Exception as exc:
        print("\nFEHLER beim Import:")
        print(exc)
        return 1

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
